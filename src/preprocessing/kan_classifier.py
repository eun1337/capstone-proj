# =====================================================================
# kan_classifier.py
# KAN 상품분류코드 자동 분류 스크립트
# - 지원 모델: Gemini / OpenAI GPT / Anthropic Claude / Clova
# - MODEL_PROVIDER만 변경하여 모델 전환 가능
# - API 키는 api_keys.env 파일에서 안전하게 로딩
# - INPUT_FILE을 변경하여 complete / missing 파일 개별 처리
# - 결과는 서식 포함 .xlsx 파일로 저장
# - 어느 디렉토리에서 실행해도 경로 자동 계산 (팀 공유 환경 대응)
# =====================================================================

import pathlib
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import time
import os
import json
import re
from dotenv import load_dotenv
from tqdm import tqdm

# =====================================================================
# 경로 기준 설정 (스크립트 위치 기준 → 어디서 실행해도 동일하게 동작)
# =====================================================================

_HERE = pathlib.Path(__file__).resolve().parent          # src/preprocessing/
_DATA = (_HERE / "../../data/master").resolve()          # data/master/

# =====================================================================
# ★ 설정 (실행 전 여기만 수정) ★
# =====================================================================

# 사용할 모델 제공자: "gemini" | "openai" | "anthropic" | "clova"
MODEL_PROVIDER = "clova"

# 각 제공자별 사용 모델명 (필요시 변경)
MODEL_NAMES = {
    "gemini":    "gemini-2.5-flash",
    "openai":    "gpt-4o-mini",
    "anthropic": "claude-sonnet-4-5",
    "clova":     "HCX-007",
}

# 입력 파일 경로 (complete 또는 missing 파일로 변경)
INPUT_FILE = str(_DATA / "a_llm_testset_375_missing.xlsx"
"")
KAN_FILE   = str(_DATA / "[대한상공회의소]KAN상품분류코드.xlsx")

# 출력 파일명: None이면 자동 생성
OUTPUT_FILE = None

# API 키 파일 경로 (스크립트와 같은 폴더: src/preprocessing/api_keys.env)
ENV_FILE = str(_HERE / "api_keys.env")

BATCH_SIZE    = 5
SLEEP_SECONDS = 15
START_IDX     = 0     # 이어서 돌릴 때 시작 행 번호 (0 = 처음부터)


# =====================================================================
# API 키 로딩 (.env 파일)
# =====================================================================

def load_api_key(provider: str, env_file: str):
    """api_keys.env 파일에서 해당 provider의 API 키를 로딩합니다."""
    if not os.path.exists(env_file):
        raise FileNotFoundError(
            f"❌ API 키 파일 없음: '{env_file}'\n"
            f"   src/preprocessing/ 폴더에 api_keys.env 파일을 만들고 아래 형식으로 입력하세요.\n"
            f"   GEMINI_API_KEY=AIza...\n"
            f"   OPENAI_API_KEY=sk-...\n"
            f"   ANTHROPIC_API_KEY=sk-ant-...\n"
            f"   CLOVA_API_KEY=nv-..."
        )
    load_dotenv(env_file)
    key_map = {
        "gemini":    "GEMINI_API_KEY",
        "openai":    "OPENAI_API_KEY",
        "anthropic": "ANTHROPIC_API_KEY",
        "clova":     "CLOVA_API_KEY",
    }
    env_var = key_map[provider]
    key = os.getenv(env_var)
    if not key or key.startswith("YOUR_"):
        raise ValueError(
            f"❌ {env_var}이 설정되지 않았습니다.\n"
            f"   '{env_file}' 파일에 {env_var}=실제키값 형태로 입력해 주세요."
        )
    return key


# =====================================================================
# KAN 분류 체계 텍스트 빌드 (소분류까지)
# =====================================================================

def build_taxonomy_text(kan_file: str) -> str:
    df = pd.read_excel(kan_file, dtype=str)
    df["KAN_CODE"] = df["KAN_CODE"].str.zfill(8)
    df["code6"] = df["KAN_CODE"].str[:6]

    lines = []
    prev_major = prev_mid = prev_minor = None
    detail_buf = []
    prev_code6 = None

    def flush(code6, minor, details):
        lines.append(f'    소분류: [{code6}] {minor} → {"│".join(details)}')

    for _, row in df.iterrows():
        major  = row["CLS_NM_1"]
        mid    = row["CLS_NM_2"]
        minor  = row["CLS_NM_3"]
        detail = row["CLS_NM_4"]
        code6  = row["code6"]

        if major != prev_major:
            if prev_minor:
                flush(prev_code6, prev_minor, detail_buf)
                detail_buf = []
            lines.append(f"대분류: {major}")
            prev_major = major
            prev_mid = prev_minor = None

        if mid != prev_mid:
            if prev_minor:
                flush(prev_code6, prev_minor, detail_buf)
                detail_buf = []
            lines.append(f"  중분류: {mid}")
            prev_mid = mid
            prev_minor = None

        if minor != prev_minor:
            if prev_minor:
                flush(prev_code6, prev_minor, detail_buf)
                detail_buf = []
            prev_minor = minor
            prev_code6 = code6

        detail_buf.append(detail)

    if detail_buf:
        flush(prev_code6, prev_minor, detail_buf)

    return "\n".join(lines)


# =====================================================================
# 시스템 프롬프트 빌드
# =====================================================================

def build_system_prompt(kan_file: str) -> str:
    taxonomy = build_taxonomy_text(kan_file)
    return f"""당신은 대한상공회의소 KAN 상품분류코드 체계에 따라 물류 상품을 분류하는 전문가입니다.

## 역할
입력된 상품 정보(바코드, 상품명)를 아래 [KAN 분류 체계]에 따라
대분류 / 중분류 / 소분류(6자리 코드 포함)로 정확하게 분류하십시오.

## KAN 코드 구조
- 전체 8자리 구조: [대분류 2자리][중분류 2자리][소분류 2자리][세분류 2자리]
- 출력 대상: 소분류까지의 앞 6자리 코드 (세분류는 출력하지 않음)
- 소분류 코드는 반드시 아래 분류 체계의 [코드] 중 하나여야 합니다

## KAN 분류 체계 (대분류 > 중분류 > 소분류[코드] → 세분류 목록)
세분류는 각 소분류에 속하는 품목 예시입니다. 상품명을 세분류 품목과 대조하여 가장 적합한 소분류를 찾으십시오.

{taxonomy}

## 입력 형식
한 번에 여러 상품이 번호 목록으로 제공됩니다.
일부 상품에는 [참고] 태그로 기존 분류 정보가 함께 제공될 수 있습니다.

예시 입력:
1. 바코드: 8801234567890 | 상품명: 농심 신라면 120g | [참고] 대분류: 식품 | 중분류: 면류 | 소분류: 라면
2. 바코드: 9900987654321 | 상품명: 테팔 프라이팬 28cm

## 출력 형식
반드시 아래 JSON 배열 형식만 출력하십시오.
설명, 이유, 추가 텍스트, 마크다운 코드블록(```json)은 절대 포함하지 마십시오.

[
  {{"번호": 1, "소분류코드": "011204", "대분류": "가공식품", "중분류": "즉석/편의식품", "소분류": "라면류"}},
  {{"번호": 2, "소분류코드": "031102", "대분류": "일상용품", "중분류": "주방용품", "소분류": "조리용기"}}
]

## 분류 규칙
1. [세분류 대조 우선] 상품명과 세분류 목록(→ 뒤의 항목들)을 먼저 대조하여 해당 소분류를 특정하십시오.
2. [코드 엄격 준수] 소분류코드는 반드시 위 분류 체계의 [코드] 목록에 있는 6자리 값 중 하나여야 합니다.
3. [본질 기반 분류] 브랜드명·용량·규격 등 부가 정보를 배제하고 상품의 본질적 속성·용도 기준으로 분류하십시오.
4. [핵심 용도 우선] 상품이 여러 분류에 걸칠 경우, 소비자의 주된 구매 목적을 기준으로 분류하십시오.
5. [기존 분류 참고] 입력에 [참고] 기존 분류가 있으면 상품 성격 파악에 가볍게 참고해도 좋습니다.
6. [최선 매핑 우선] 세분류 목록에 정확히 일치하지 않아도, 본질적 속성이 가장 유사한 소분류로 매핑하십시오.
7. [완전 불가 시] 전체 분류 체계 어디에도 본질적으로 맞지 않으면 소분류코드: "999999", 소분류: "기타"로 표기하십시오.
8. [바코드 보조 활용] 바코드는 상품 식별 보조 정보로만 참고하고, 최종 판단은 상품명 기준으로 하십시오.
"""


# =====================================================================
# 유효 코드 목록 로딩 (검증용)
# =====================================================================

def load_valid_codes(kan_file: str) -> set:
    df = pd.read_excel(kan_file, dtype={"KAN_CODE": str})
    df["KAN_CODE"] = df["KAN_CODE"].str.zfill(8)
    return set(df["KAN_CODE"].str[:6].unique())


# =====================================================================
# 데이터 로딩 및 records 구성
# =====================================================================

def row_to_record(i: int, row: pd.Series) -> dict:
    def safe(col):
        v = row.get(col, "")
        return "" if pd.isna(v) else str(v).strip()

    return {
        "idx":       i + 1,
        "sample_id": safe("sample_id"),
        "test_case": safe("테스트케이스"),
        "barcode":   safe("바코드"),
        "name":      safe("상품명"),
        "cls1":      safe("기존_대분류"),
        "cls2":      safe("기존_중분류"),
        "cls3":      safe("기존_소분류"),
    }


# =====================================================================
# 배치 유저 프롬프트 빌드
# =====================================================================

def build_batch_prompt(batch: list) -> str:
    lines = []
    for item in batch:
        parts = [f"{item['idx']}. 바코드: {item['barcode'] or 'N/A'} | 상품명: {item['name']}"]
        ref_parts = []
        if item.get("cls1"):
            ref_parts.append(f"대분류: {item['cls1']}")
        if item.get("cls2"):
            ref_parts.append(f"중분류: {item['cls2']}")
        if item.get("cls3"):
            ref_parts.append(f"소분류: {item['cls3']}")
        if ref_parts:
            parts.append(f"[참고] {' | '.join(ref_parts)}")
        lines.append(" | ".join(parts))
    return "\n".join(lines)


# =====================================================================
# JSON 응답 파싱
# =====================================================================

def parse_response(response_text: str, batch: list) -> list:
    cleaned = re.sub(r"```(?:json)?", "", response_text).strip().rstrip("`")
    try:
        results = json.loads(cleaned)
        # 번호 키 없는 경우 보정
        for i, r in enumerate(results):
            if "번호" not in r:
                r["번호"] = batch[i]["idx"]
        return results
    except json.JSONDecodeError:
        return [
            {
                "번호": item["idx"],
                "소분류코드": "PARSE_ERROR",
                "대분류": "파싱실패",
                "중분류": "파싱실패",
                "소분류": response_text[:80],
            }
            for item in batch
        ]


# =====================================================================
# 모델별 API 클라이언트 초기화
# =====================================================================

def init_client(provider: str, api_key, model_name: str, system_prompt: str):
    if provider == "gemini":
        import google.generativeai as genai
        genai.configure(api_key=api_key)
        client = genai.GenerativeModel(
            model_name=model_name,
            system_instruction=system_prompt,
        )

    elif provider == "openai":
        from openai import OpenAI
        client = OpenAI(api_key=api_key)
        client._kan_system_prompt = system_prompt
        client._kan_model_name    = model_name

    elif provider == "anthropic":
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)
        client._kan_system_prompt = system_prompt
        client._kan_model_name    = model_name

    elif provider == "clova":
        import types
        client = types.SimpleNamespace()
        client._kan_system_prompt = system_prompt
        client._kan_model_name    = model_name
        client._kan_api_key       = api_key

    else:
        raise ValueError(f"지원하지 않는 MODEL_PROVIDER: '{provider}'  →  gemini | openai | anthropic | clova 중 선택")

    return client


# =====================================================================
# 모델별 단일 배치 호출
# =====================================================================

def _call_gemini(client, prompt: str) -> str:
    response = client.generate_content(
        prompt,
        generation_config={"temperature": 0},
    )
    if not response.text:
        raise ValueError("빈 응답 (Gemini)")
    return response.text.strip()


def _call_openai(client, prompt: str) -> str:
    response = client.chat.completions.create(
        model=client._kan_model_name,
        temperature=0,
        messages=[
            {"role": "system", "content": client._kan_system_prompt},
            {"role": "user",   "content": prompt},
        ],
    )
    text = response.choices[0].message.content
    if not text:
        raise ValueError("빈 응답 (OpenAI)")
    return text.strip()


def _call_anthropic(client, prompt: str) -> str:
    response = client.messages.create(
        model=client._kan_model_name,
        max_tokens=4096,
        temperature=0,
        system=client._kan_system_prompt,
        messages=[{"role": "user", "content": prompt}],
    )
    text = response.content[0].text
    if not text:
        raise ValueError("빈 응답 (Anthropic)")
    return text.strip()


def _call_clova(client, prompt: str) -> str:
    import requests as req
    url = f"https://clovastudio.stream.ntruss.com/v3/chat-completions/{client._kan_model_name}"
    headers = {
        "Authorization": f"Bearer {client._kan_api_key}",
        "Content-Type":  "application/json",
        "Accept":        "application/json",
    }
    body = {
        "messages": [
            {"role": "system", "content": client._kan_system_prompt},
            {"role": "user",   "content": prompt},
        ],
        "maxCompletionTokens": 4096, # CLOVA 아닌경우, "maxTokens":   4096,
        "temperature": 0.0,
        "topP":        0.8,
        "thinking":    {"effort": "none"},  # 추론 모드 끄기 (속도/토큰 절약)
    }
    response = req.post(url, headers=headers, json=body, timeout=60)

    if response.status_code != 200:
        raise ValueError(f"Clova API 오류 [{response.status_code}]: {response.text[:200]}")

    result = response.json()

    # 응답 구조: result.message.content
    try:
        text = result["result"]["message"]["content"]
    except (KeyError, TypeError):
        raise ValueError(f"Clova 응답 파싱 실패: {str(result)[:200]}")

    if not text:
        raise ValueError("빈 응답 (Clova)")
    return text.strip()


_CALLERS = {
    "gemini":    _call_gemini,
    "openai":    _call_openai,
    "anthropic": _call_anthropic,
    "clova":     _call_clova,
}


# =====================================================================
# 배치 분류 (재시도 포함)
# =====================================================================

def classify_batch(provider: str, client, batch: list, retry: int = 0) -> list:
    try:
        prompt   = build_batch_prompt(batch)
        raw_text = _CALLERS[provider](client, prompt)
        return parse_response(raw_text, batch)

    except Exception as e:
        err_str = str(e)
        is_rate_limit = any(code in err_str for code in ["429", "rate_limit", "RateLimitError", "ResourceExhausted"])

        if is_rate_limit and retry < 3:
            wait = 15 * (retry + 1)
            print(f"\n  [속도 제한] {wait}초 대기 후 재시도 (retry={retry + 1})...")
            time.sleep(wait)
            return classify_batch(provider, client, batch, retry + 1)

        return [
            {
                "번호": item["idx"],
                "소분류코드": "API_ERROR",
                "대분류": "API오류",
                "중분류": "API오류",
                "소분류": err_str[:80],
            }
            for item in batch
        ]


# =====================================================================
# 출력 파일명 자동 생성
# =====================================================================

def resolve_output_path(input_file: str, provider: str) -> str:
    base = os.path.splitext(input_file)[0]
    return f"{base}_{provider}_result.xlsx"


# =====================================================================
# 엑셀 저장 (서식 포함)
# =====================================================================

def save_xlsx(df: pd.DataFrame, output_path: str, provider: str) -> None:
    kan_cols  = ["kan_소분류코드", "kan_대분류", "kan_중분류", "kan_소분류", "코드검증"]
    orig_cols = [c for c in df.columns if c not in kan_cols]
    df = df[orig_cols + kan_cols]

    df.to_excel(output_path, index=False, engine="openpyxl")

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    COLOR_HEADER_ORIG = "4472C4"
    COLOR_HEADER_KAN  = "ED7D31"
    COLOR_OK          = "E2EFDA"
    COLOR_ETC         = "FFF2CC"
    COLOR_ERR         = "FCE4D6"

    fill_header_orig = PatternFill("solid", fgColor=COLOR_HEADER_ORIG)
    fill_header_kan  = PatternFill("solid", fgColor=COLOR_HEADER_KAN)
    fill_ok          = PatternFill("solid", fgColor=COLOR_OK)
    fill_etc         = PatternFill("solid", fgColor=COLOR_ETC)
    fill_err         = PatternFill("solid", fgColor=COLOR_ERR)

    font_header = Font(bold=True, color="FFFFFF", size=10)
    font_normal = Font(size=10)

    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=False)

    kan_col_names = set(kan_cols)
    col_index_map = {cell.value: cell.column for cell in ws[1]}

    for cell in ws[1]:
        cell.fill      = fill_header_kan if cell.value in kan_col_names else fill_header_orig
        cell.font      = font_header
        cell.alignment = center
        cell.border    = border
    ws.row_dimensions[1].height = 22

    verify_col      = col_index_map.get("코드검증")
    kan_col_indices = {col_index_map[c] for c in kan_cols if c in col_index_map}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        verify_val = row[verify_col - 1].value if verify_col else None
        if   verify_val == "정상":          row_fill = fill_ok
        elif verify_val == "기타(999999)":  row_fill = fill_etc
        elif verify_val == "코드오류":      row_fill = fill_err
        else:                               row_fill = None

        for cell in row:
            cell.font   = font_normal
            cell.border = border
            if row_fill and cell.column in kan_col_indices:
                cell.fill = row_fill
            header_name = ws.cell(1, cell.column).value
            cell.alignment = left if header_name == "상품명" else center

    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max(
            len(str(col_cells[0].value or "")),
            max((len(str(c.value or "")) for c in col_cells[1:]), default=0)
        )
        ws.column_dimensions[col_letter].width = min(max_len * 1.3 + 2, 45)

    ws.freeze_panes = "A2"
    ws.title = f"KAN분류_{provider}"

    wb.save(output_path)


# =====================================================================
# 메인
# =====================================================================

def main():
    for f in [INPUT_FILE, KAN_FILE]:
        if not os.path.exists(f):
            print(f"❌ 파일 없음: '{f}'")
            return

    try:
        api_key = load_api_key(MODEL_PROVIDER, ENV_FILE)
        print(f"🔑 API 키 로딩 완료 ({MODEL_PROVIDER})")
    except (FileNotFoundError, ValueError) as e:
        print(e)
        return

    output_path = OUTPUT_FILE if OUTPUT_FILE else resolve_output_path(INPUT_FILE, MODEL_PROVIDER)

    print("📂 KAN 유효 코드 로딩...")
    valid_codes = load_valid_codes(KAN_FILE)
    print(f"  → 소분류 코드 {len(valid_codes)}개")

    print("📂 시스템 프롬프트 빌드...")
    system_prompt = build_system_prompt(KAN_FILE)

    model_name = MODEL_NAMES[MODEL_PROVIDER]
    print(f"🤖 모델 초기화: [{MODEL_PROVIDER}] {model_name}")
    client = init_client(MODEL_PROVIDER, api_key, model_name, system_prompt)
    print(f"  → 준비 완료")

    df_full = pd.read_excel(INPUT_FILE, dtype={"바코드": str})
    df      = df_full.iloc[START_IDX:].reset_index(drop=True)
    print(f"📊 처리 대상: {START_IDX} ~ {START_IDX + len(df) - 1}행 ({len(df):,}개)")

    records = [row_to_record(i, row) for i, row in df.iterrows()]
    batches = [records[i:i + BATCH_SIZE] for i in range(0, len(records), BATCH_SIZE)]

    all_results: dict = {}
    print(f"\n🚀 분류 시작 (배치 {len(batches)}개, 배치 크기 {BATCH_SIZE})...")

    for batch in tqdm(batches, desc="분류 진행"):
        for r in classify_batch(MODEL_PROVIDER, client, batch):
            all_results[r["번호"]] = r
        if len(batches) > 1:
            time.sleep(SLEEP_SECONDS)

    df["kan_소분류코드"] = [all_results.get(i + 1, {}).get("소분류코드", "ERROR") for i in range(len(df))]
    df["kan_대분류"]     = [all_results.get(i + 1, {}).get("대분류",    "ERROR") for i in range(len(df))]
    df["kan_중분류"]     = [all_results.get(i + 1, {}).get("중분류",    "ERROR") for i in range(len(df))]
    df["kan_소분류"]     = [all_results.get(i + 1, {}).get("소분류",    "ERROR") for i in range(len(df))]
    df["코드검증"] = df["kan_소분류코드"].apply(
        lambda c: "정상" if c in valid_codes else ("기타(999999)" if c == "999999" else "코드오류")
    )

    save_xlsx(df, output_path, MODEL_PROVIDER)

    total = len(df)
    ok    = (df["코드검증"] == "정상").sum()
    etc   = (df["코드검증"] == "기타(999999)").sum()
    err   = (df["코드검증"] == "코드오류").sum()
    print(f"\n✅ 완료 → '{output_path}'")
    print(f"   정상: {ok:,} | 기타(999999): {etc:,} | 코드오류: {err:,} / 전체: {total:,}")
    print(f"   정확도(정상/전체): {ok / total * 100:.1f}%")


if __name__ == "__main__":
    main()