"""
evaluate_kan_classifier.py

설명 : LLM 모델별 KAN 분류 결과를 정답과 비교하여 채점합니다.
    대/중/소분류 명칭 텍스트 완전 일치 기준으로 정오답을 판정하고,
    모델별 정확도 요약과 행별 정오답 결과를 .xlsx 파일로 저장합니다.

"""

import pathlib
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter



_HERE = pathlib.Path(__file__).resolve().parent
_DATA = (_HERE / "../../data/master").resolve()



INPUT_FILE  = str(_DATA / "a_llm_testset_375_final.xlsx")
OUTPUT_FILE = str(_DATA / "a_llm_testset_375_final_scored.xlsx")


MODELS = [
    ("GPT",   "GPT_대분류",   "GPT_중분류",   "GPT_소분류"),
    ("Gemini","Gemini_대분류","Gemini_중분류","Gemini_소분류"),
    ("Claude","Claude_대분류","Claude_중분류","Claude_소분류"),
    ("Clova", "Clova_대분류", "Clova_중분류", "Clova_소분류"),
]

# 정답 컬럼명
ANS_CLS1 = "정답_대분류"
ANS_CLS2 = "정답_중분류"
ANS_CLS3 = "정답_소분류"

# 상품명 컬럼명
NAME_COL  = "상품명"


def score_row(row, model_cls1, model_cls2, model_cls3):
    """대/중/소 모두 일치하면 정답(1), 아니면 오답(0)."""
    ans1 = str(row[ANS_CLS1]).strip()
    ans2 = str(row[ANS_CLS2]).strip()
    ans3 = str(row[ANS_CLS3]).strip()
    pred1 = str(row[model_cls1]).strip()
    pred2 = str(row[model_cls2]).strip()
    pred3 = str(row[model_cls3]).strip()
    return 1 if (ans1 == pred1 and ans2 == pred2 and ans3 == pred3) else 0



def save_xlsx(df: pd.DataFrame, summary: pd.DataFrame, output_path: str) -> None:

    # -- 색상 ----------------------------------------------------------
    COLOR_HEADER    = "4472C4"
    COLOR_SUMMARY   = "ED7D31"
    COLOR_CORRECT   = "E2EFDA"   # 연초록 - 정답
    COLOR_WRONG     = "FCE4D6"   # 연빨강 - 오답
    COLOR_NAME      = "F2F2F2"   # 연회색 - 상품명/정답 영역

    fill_header  = PatternFill("solid", fgColor=COLOR_HEADER)
    fill_summary = PatternFill("solid", fgColor=COLOR_SUMMARY)
    fill_correct = PatternFill("solid", fgColor=COLOR_CORRECT)
    fill_wrong   = PatternFill("solid", fgColor=COLOR_WRONG)
    fill_name    = PatternFill("solid", fgColor=COLOR_NAME)

    font_header  = Font(bold=True, color="FFFFFF", size=10)
    font_bold    = Font(bold=True, size=10)
    font_normal  = Font(size=10)

    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left   = Alignment(horizontal="left",   vertical="center")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="채점결과", index=False)
        summary.to_excel(writer, sheet_name="모델별요약", index=False)

    wb = openpyxl.load_workbook(output_path)

    ws = wb["채점결과"]
    col_index_map = {cell.value: cell.column for cell in ws[1]}


    score_cols = {col_index_map[f"{m}_정오답"] for m, *_ in MODELS if f"{m}_정오답" in col_index_map}
    name_cols  = {col_index_map[c] for c in [NAME_COL, ANS_CLS1, ANS_CLS2, ANS_CLS3] if c in col_index_map}


    for cell in ws[1]:
        col_name = cell.value or ""
        if col_name in [NAME_COL, ANS_CLS1, ANS_CLS2, ANS_CLS3]:
            cell.fill = fill_summary
        else:
            cell.fill = fill_header
        cell.font      = font_header
        cell.alignment = center
        cell.border    = border
    ws.row_dimensions[1].height = 22


    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font   = font_normal
            cell.border = border

            if cell.column in name_cols:
                cell.fill = fill_name

            if cell.column in score_cols:
                if cell.value == "정답":
                    cell.fill = fill_correct
                    cell.font = Font(bold=True, size=10, color="375623")
                elif cell.value == "오답":
                    cell.fill = fill_wrong
                    cell.font = Font(bold=True, size=10, color="9C0006")

            header_name = ws.cell(1, cell.column).value
            cell.alignment = left if header_name == NAME_COL else center

    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max(
            len(str(col_cells[0].value or "")),
            max((len(str(c.value or "")) for c in col_cells[1:]), default=0)
        )
        ws.column_dimensions[col_letter].width = min(max_len * 1.3 + 2, 40)

    ws.freeze_panes = "A2"


    ws2 = wb["모델별요약"]

    for cell in ws2[1]:
        cell.fill      = fill_header
        cell.font      = font_header
        cell.alignment = center
        cell.border    = border
    ws2.row_dimensions[1].height = 22

    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        for cell in row:
            cell.font      = font_bold if cell.column == 1 else font_normal
            cell.alignment = center
            cell.border    = border
   
            header_name = ws2.cell(1, cell.column).value
            if header_name == "정확도(%)" and cell.value is not None:
                val = float(cell.value)
                if val >= 80:
                    cell.fill = fill_correct
                elif val >= 60:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
                else:
                    cell.fill = fill_wrong

    for col_cells in ws2.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max(
            len(str(col_cells[0].value or "")),
            max((len(str(c.value or "")) for c in col_cells[1:]), default=0)
        )
        ws2.column_dimensions[col_letter].width = max_len * 1.5 + 4

    ws2.freeze_panes = "A2"

    wb.save(output_path)

def main():
    import os
    if not os.path.exists(INPUT_FILE):
        print(f" 파일 없음: '{INPUT_FILE}'")
        return

    print("📂 파일 로딩...")
    df = pd.read_excel(INPUT_FILE, dtype=str)
    print(f"  → {len(df):,}행 로딩 완료")

    print("📝 채점 중...")
    result_cols = []
    summary_rows = []

    for model_name, col1, col2, col3 in MODELS:
        missing = [c for c in [col1, col2, col3] if c not in df.columns]
        if missing:
            print(f"    {model_name}: 컬럼 없음 {missing} → 건너뜀")
            continue

        score_col = f"{model_name}_정오답"
        df[score_col] = df.apply(
            lambda row: "정답" if score_row(row, col1, col2, col3) == 1 else "오답",
            axis=1
        )
        result_cols.append(score_col)

        total   = len(df)
        correct = (df[score_col] == "정답").sum()
        acc     = round(correct / total * 100, 1)
        summary_rows.append({
            "모델":     model_name,
            "전체":     total,
            "정답":     correct,
            "오답":     total - correct,
            "정확도(%)": acc,
        })
        print(f"   {model_name}: {correct}/{total} ({acc}%)")

    base_cols    = [NAME_COL, ANS_CLS1, ANS_CLS2, ANS_CLS3]
    base_cols    = [c for c in base_cols if c in df.columns]
    model_blocks = []
    for model_name, col1, col2, col3 in MODELS:
        block = [c for c in [col1, col2, col3, f"{model_name}_정오답"] if c in df.columns]
        model_blocks.extend(block)
    other_cols = [c for c in df.columns if c not in base_cols + model_blocks]
    final_cols = base_cols + model_blocks

    df_out = df[final_cols]

    summary_df = pd.DataFrame(summary_rows)

    print(f"\n💾 저장 중: '{OUTPUT_FILE}'")
    save_xlsx(df_out, summary_df, OUTPUT_FILE)

    print(f"\n완료 → '{OUTPUT_FILE}'")
    print("\n모델별 정확도 요약")
    print("-" * 40)
    for row in summary_rows:
        print(f"  {row['모델']:10s}: {row['정답']:3d}/{row['전체']} ({row['정확도(%)']:.1f}%)")


if __name__ == "__main__":
    main()