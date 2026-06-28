# =====================================================================
# classify_ab_master_gemini.py
# ab_final_product_master_2.xlsx 전용 Gemini 기반 KAN 분류 "구간별" 배치 스크립트
# (대용량/비용방어형 리팩토링 버전)
#
# - kan_classifier.py의 시스템 프롬프트 / 프롬프트 빌더 / 파서를 재사용
#   (MODEL_PROVIDER는 gemini로 고정)
# - 지정한 행 구간(start_row~end_row, 1-based, 헤더 제외)만 처리
# - 처리 후 해당 구간의 셀만 정확히 덮어쓰기 (다른 행은 절대 건드리지 않음)
# - 실행 전 원본 백업(ab_final_product_master_2_backup.xlsx)을 최초 1회만 생성
#   (이미 백업이 있으면 보존을 위해 다시 덮어쓰지 않음)
# - 소분류코드는 항상 6자리 문자열(zfill)로 강제하고, 셀 서식을 텍스트('@')로
#   지정하여 엑셀에서 앞자리 0이 사라지지 않도록 함
# - 출력 컬럼: Gemini_대분류 / Gemini_중분류 / Gemini_소분류 / KAN_CODE(소분류코드 6자리)
# - 배치(20건) 처리 완료마다 즉시 엑셀에 물리적 저장(체크포인트) → 중간에 에러가
#   나도 그 시점까지 처리된 결과는 보존됨
# - 429(ResourceExhausted) 발생 시 60초 단위로 대기 시간을 늘려가며 재시도
#
# 사용법:
#   python classify_ab_master_gemini.py <시작행> <종료행>
#   예) python classify_ab_master_gemini.py 1 10     → 선제 테스트 (10건)
#       python classify_ab_master_gemini.py 1 800    → 본 주행 1구간
# =====================================================================

import argparse
import os
import pathlib
import random
import shutil
import sys
import time

import openpyxl
from tqdm import tqdm

# =====================================================================
# 경로 기준 설정 (스크립트 위치 기준 → 어디서 실행해도 동일하게 동작)
# =====================================================================

_HERE = pathlib.Path(__file__).resolve().parent          # src/preprocessing/
_DATA = (_HERE / "../../data/master").resolve()          # data/master/

sys.path.insert(0, str(_HERE))
from kan_classifier import (  # noqa: E402  (경로 설정 이후 import 필요)
    build_batch_prompt,
    build_system_prompt,
    init_client,
    load_api_key,
    load_valid_codes,
    parse_response,
)

# =====================================================================
# ★ 설정 (실행 전 여기만 수정) ★
# =====================================================================

TARGET_FILE = _DATA / "ab_final_product_master_2.xlsx"
BACKUP_FILE = _DATA / "ab_final_product_master_2_backup.xlsx"
KAN_FILE    = _DATA / "[대한상공회의소]KAN상품분류코드.xlsx"
ENV_FILE    = _HERE / "api_keys.env"

MODEL_PROVIDER = "gemini"
MODEL_NAME     = "gemini-2.5-flash"

BATCH_SIZE = 20   # 1회 API 호출당 상품 수
# ※ 시스템 프롬프트(KAN 분류체계 전체, 약 2만 토큰)가 호출마다 매번 재전송되므로
#   BATCH_SIZE를 5→20으로 늘려 같은 처리량 대비 호출 수(=누적 토큰 소비)를 1/4로 줄임.

SLEEP_SECONDS_MIN = 20   # 배치 사이 대기 시간(초) 하한
SLEEP_SECONDS_MAX = 30   # 배치 사이 대기 시간(초) 상한 → 매번 이 구간에서 랜덤 대기

CHECKPOINT_EVERY_BATCHES = 1   # N배치(=N*BATCH_SIZE건)마다 즉시 엑셀에 중간 저장
#   BATCH_SIZE=20 기준 1배치=20건이므로, 1로 두면 "20건마다 저장" 요건과 동일

RATE_LIMIT_BASE_WAIT  = 60   # 429(ResourceExhausted) 발생 시 최소 대기(초). 재시도마다 배수로 증가
RATE_LIMIT_MAX_RETRY  = 5
OTHER_ERROR_MAX_RETRY = 2    # 429가 아닌 일시적 오류에 대한 재시도 횟수(대기는 짧게)

# 입력 컬럼명 (ab_final_product_master_2.xlsx 기준)
COL_BARCODE = "바코드"
COL_NAME    = "상품명"
COL_CLS1    = "대분류"
COL_CLS2    = "중분류"
COL_CLS3    = "소분류"

# 출력 컬럼명 (Gemini 분류 결과 기록 대상)
COL_OUT_CLS1 = "Gemini_대분류"
COL_OUT_CLS2 = "Gemini_중분류"
COL_OUT_CLS3 = "Gemini_소분류"
COL_OUT_CODE = "KAN_CODE"   # 기존의 빈 KAN_CODE 컬럼에 소분류코드 6자리를 기록

REQUIRED_COLUMNS = [
    COL_BARCODE, COL_NAME, COL_CLS1, COL_CLS2, COL_CLS3,
    COL_OUT_CLS1, COL_OUT_CLS2, COL_OUT_CLS3, COL_OUT_CODE,
]


# =====================================================================
# 안전장치
# =====================================================================

def is_excel_file_open(path: pathlib.Path) -> bool:
    """엑셀이 파일을 열고 있으면 같은 폴더에 '~$파일명' 잠금 파일이 생성된다."""
    lock_path = path.parent / f"~${path.name}"
    return lock_path.exists()


def ensure_backup(target: pathlib.Path, backup: pathlib.Path) -> None:
    if backup.exists():
        print(f"백업 파일 이미 존재 → 원본 보존을 위해 다시 만들지 않음: '{backup.name}'")
        return
    shutil.copy2(target, backup)
    print(f"원본 백업 생성 완료: '{backup.name}'")


def normalize_code(raw) -> str:
    """숫자로만 구성된 코드만 6자리로 zfill. 오류 마커(API_ERROR 등)는 그대로 둔다."""
    s = str(raw).strip()
    return s.zfill(6) if s.isdigit() else s


def map_columns(ws) -> dict:
    header = {cell.value: cell.column for cell in ws[1]}
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ValueError(f"필수 컬럼 누락: {missing}")
    return header


def safe_cell(ws, row: int, col: int) -> str:
    v = ws.cell(row=row, column=col).value
    return "" if v is None else str(v).strip()


def build_records(ws, col_map: dict, start_row: int, end_row: int) -> list:
    """start_row~end_row(1-based, 헤더 제외)에 대응하는 records 구성.
    엑셀 시트상 실제 행 번호 = 데이터 행 번호 + 1 (1행은 헤더)."""
    records = []
    for i, excel_row in enumerate(range(start_row + 1, end_row + 2)):
        records.append({
            "idx":     i + 1,
            "barcode": safe_cell(ws, excel_row, col_map[COL_BARCODE]),
            "name":    safe_cell(ws, excel_row, col_map[COL_NAME]),
            "cls1":    safe_cell(ws, excel_row, col_map[COL_CLS1]),
            "cls2":    safe_cell(ws, excel_row, col_map[COL_CLS2]),
            "cls3":    safe_cell(ws, excel_row, col_map[COL_CLS3]),
        })
    return records



# Gemini 호출 (429는 60초+ 장기 백오프, 그 외 오류는 짧게 재시도)

def _is_rate_limit(err_str: str) -> bool:
    return any(code in err_str for code in
               ["429", "ResourceExhausted", "rate_limit", "RateLimitError", "quota"])


def call_gemini(client, prompt: str, retry: int = 0, rate_limit_retry: int = 0) -> str:
    try:
        response = client.generate_content(prompt, generation_config={"temperature": 0})
        if not response.text:
            raise ValueError("빈 응답 (Gemini)")
        return response.text.strip()

    except Exception as e:
        err_str = str(e)

        if _is_rate_limit(err_str) and rate_limit_retry < RATE_LIMIT_MAX_RETRY:
            wait = RATE_LIMIT_BASE_WAIT * (rate_limit_retry + 1)
            print(f"\n  ⏳ [429 Rate Limit] {wait}초 대기 후 재시도 "
                  f"({rate_limit_retry + 1}/{RATE_LIMIT_MAX_RETRY})...")
            time.sleep(wait)
            return call_gemini(client, prompt, retry, rate_limit_retry + 1)

        if not _is_rate_limit(err_str) and retry < OTHER_ERROR_MAX_RETRY:
            wait = 15 * (retry + 1)
            print(f"\n  ⚠️  [오류] {err_str[:120]} → {wait}초 후 재시도 "
                  f"({retry + 1}/{OTHER_ERROR_MAX_RETRY})...")
            time.sleep(wait)
            return call_gemini(client, prompt, retry + 1, rate_limit_retry)

        raise


def classify_batch_gemini(client, batch: list) -> list:
    prompt = build_batch_prompt(batch)
    try:
        raw_text = call_gemini(client, prompt)
        return parse_response(raw_text, batch)
    except Exception as e:
        err_str = str(e)
        return [
            {
                "번호": item["idx"],
                "소분류코드": "API_ERROR",
                "대분류": "API오류",
                "중분류": "API오류",
                "소분류": err_str[:80],
            }
            for item in batch
        ]


# =====================================================================
# 메인
# =====================================================================

def main(start_row: int, end_row: int) -> None:
    for f in [TARGET_FILE, KAN_FILE]:
        if not f.exists():
            print(f"❌ 파일 없음: '{f}'")
            return

    if is_excel_file_open(TARGET_FILE):
        print(f"❌ '{TARGET_FILE.name}' 파일이 엑셀에서 열려 있는 것으로 보입니다 (잠금 파일 발견).")
        print("   엑셀을 완전히 닫고 다시 실행해 주세요. (열어둔 채로 저장하면 실패하거나 충돌할 수 있습니다)")
        return

    if start_row < 1 or end_row < start_row:
        print(f"❌ 잘못된 행 범위: {start_row} ~ {end_row}")
        return

    print(f"📂 엑셀 로딩: '{TARGET_FILE.name}'")
    wb = openpyxl.load_workbook(TARGET_FILE)
    ws = wb.active

    try:
        col_map = map_columns(ws)
    except ValueError as e:
        print(f"❌ {e}")
        return

    total_data_rows = ws.max_row - 1
    if end_row > total_data_rows:
        print(f"❌ 종료 행({end_row})이 전체 데이터 행 수({total_data_rows})를 초과합니다.")
        return

    ensure_backup(TARGET_FILE, BACKUP_FILE)

    try:
        api_key = load_api_key(MODEL_PROVIDER, str(ENV_FILE))
        print(f"🔑 API 키 로딩 완료 ({MODEL_PROVIDER})")
    except (FileNotFoundError, ValueError) as e:
        print(e)
        return

    print("📂 KAN 유효 코드 / 시스템 프롬프트 빌드...")
    valid_codes   = load_valid_codes(str(KAN_FILE))
    system_prompt = build_system_prompt(str(KAN_FILE))
    print(f"  → 소분류 코드 {len(valid_codes)}개")

    print(f"🤖 모델 초기화: [{MODEL_PROVIDER}] {MODEL_NAME}")
    client = init_client(MODEL_PROVIDER, api_key, MODEL_NAME, system_prompt)

    already_done = sum(
        1 for r in range(start_row + 1, end_row + 2)
        if ws.cell(row=r, column=col_map[COL_OUT_CODE]).value not in (None, "")
    )
    if already_done:
        print(f"⚠️  이미 분류된 행 {already_done}개가 이 구간에 포함되어 있습니다 → 덮어쓰기됩니다.")

    print(f"📊 처리 대상: 데이터 {start_row}~{end_row}행 "
          f"(엑셀 시트 {start_row + 1}~{end_row + 1}행, {end_row - start_row + 1:,}개)")

    records = build_records(ws, col_map, start_row, end_row)
    batches = [records[i:i + BATCH_SIZE] for i in range(0, len(records), BATCH_SIZE)]

    counts = {"정상": 0, "기타(999999)": 0, "코드오류": 0}

    def save_checkpoint():
        # 임시 파일에 먼저 저장 후 교체 → 저장 중 오류가 나도 원본 파일은 손상되지 않음
        tmp_path = TARGET_FILE.with_suffix(".tmp.xlsx")
        wb.save(tmp_path)
        os.replace(tmp_path, TARGET_FILE)

    print(f"\n🚀 분류 시작 (배치 {len(batches)}개, 배치 크기 {BATCH_SIZE}, "
          f"체크포인트 {CHECKPOINT_EVERY_BATCHES}배치마다 즉시 저장)...")

    try:
        for batch_no, batch in enumerate(tqdm(batches, desc="분류 진행"), start=1):
            results    = classify_batch_gemini(client, batch)
            result_map = {r["번호"]: r for r in results}

            for item in batch:
                excel_row = start_row + item["idx"]
                r    = result_map.get(item["idx"], {})
                code = normalize_code(r.get("소분류코드", "ERROR"))

                ws.cell(row=excel_row, column=col_map[COL_OUT_CLS1]).value = r.get("대분류", "ERROR")
                ws.cell(row=excel_row, column=col_map[COL_OUT_CLS2]).value = r.get("중분류", "ERROR")
                ws.cell(row=excel_row, column=col_map[COL_OUT_CLS3]).value = r.get("소분류", "ERROR")

                code_cell = ws.cell(row=excel_row, column=col_map[COL_OUT_CODE])
                code_cell.value = code
                code_cell.number_format = "@"   # 텍스트 서식 강제 → 앞자리 0 보존

                if code in valid_codes:
                    counts["정상"] += 1
                elif code == "999999":
                    counts["기타(999999)"] += 1
                else:
                    counts["코드오류"] += 1

            is_last = batch_no == len(batches)
            if batch_no % CHECKPOINT_EVERY_BATCHES == 0 or is_last:
                save_checkpoint()
                done = min(batch_no * BATCH_SIZE, len(records))
                tqdm.write(f"  💾 체크포인트 저장 완료 (배치 {batch_no}/{len(batches)}, 누적 {done:,}건)")

            if not is_last:
                time.sleep(random.uniform(SLEEP_SECONDS_MIN, SLEEP_SECONDS_MAX))

    except (Exception, KeyboardInterrupt):
        print("\n🛑 처리 중 중단/오류 발생 → 지금까지 처리된 결과를 즉시 저장합니다...")
        save_checkpoint()
        print(f"   💾 임시 저장 완료 → '{TARGET_FILE.name}'")
        raise

    total = end_row - start_row + 1
    print(f"\n✅ 완료 → '{TARGET_FILE.name}' (데이터 {start_row}~{end_row}행 덮어쓰기 저장)")
    print(f"   정상: {counts['정상']:,} | 기타(999999): {counts['기타(999999)']:,} "
          f"| 코드오류: {counts['코드오류']:,} / 전체: {total:,}")
    print(f"   정확도(정상/전체): {counts['정상'] / total * 100:.1f}%")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="ab_final_product_master_2.xlsx 구간별 Gemini KAN 분류 (대용량 비용방어형)"
    )
    parser.add_argument("start_row", type=int, help="시작 행 번호 (1-based, 헤더 제외 데이터 기준)")
    parser.add_argument("end_row",   type=int, help="종료 행 번호 (1-based, inclusive)")
    args = parser.parse_args()
    main(args.start_row, args.end_row)
