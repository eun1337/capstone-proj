"""
evaluate_all_models.py

6개 파일(A/B 센터 × 1~3차)에 걸쳐 GPT·Gemini·Claude·Clova 4개 모델의
성능을 4가지 지표로 비교하고 종합 채점 보고서를 엑셀로 저장합니다.

[채점 지표]
1. 코드 오류 수 / 비율 : 모델이 예측한 소분류코드가 KAN 유효 코드 목록에 없는 경우
2. 전체 정답률         : 대/중/소분류 명칭 텍스트가 정답과 완전 일치
3. 결측 정답률         : 테스트케이스 == '결측' 행만 필터 후 정답률
4. 정상 정답률         : 테스트케이스 == '정상' 행만 필터 후 정답률
"""

import pathlib
import pandas as pd
import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

# ── 경로 설정 ────────────────────────────────────────────────────────────────
_HERE = pathlib.Path(__file__).resolve().parent
_DATA = (_HERE / "../../data/master").resolve()

KAN_FILE    = str(_DATA / "[대한상공회의소]KAN상품분류코드.xlsx")
OUTPUT_FILE = str(_DATA / "total_model_evaluation.xlsx")

# ── 입력 파일 정의 ───────────────────────────────────────────────────────────
INPUT_FILES = [
    ("A", 1, "a_llm_testset_375_final.xlsx"),
    ("A", 2, "a_llm_testset_375_final_2.xlsx"),
    ("A", 3, "a_llm_testset_375_final_3.xlsx"),
    ("B", 1, "b_llm_testset_374_final.xlsx"),
    ("B", 2, "b_llm_testset_374_final_2.xlsx"),
    ("B", 3, "b_llm_testset_374_final_3.xlsx"),
]

# ── 모델 컬럼 정의 ───────────────────────────────────────────────────────────
MODELS = [
    ("GPT",    "GPT_소분류코드",    "GPT_대분류",    "GPT_중분류",    "GPT_소분류"),
    ("Gemini", "Gemini_소분류코드", "Gemini_대분류", "Gemini_중분류", "Gemini_소분류"),
    ("Claude", "Claude_소분류코드", "Claude_대분류", "Claude_중분류", "Claude_소분류"),
    ("Clova",  "Clova_소분류코드",  "Clova_대분류",  "Clova_중분류",  "Clova_소분류"),
]

# ── 정답 / 분류 컬럼 ─────────────────────────────────────────────────────────
ANS_CLS1   = "정답_대분류"
ANS_CLS2   = "정답_중분류"
ANS_CLS3   = "정답_소분류"
TEST_CASE  = "테스트케이스"
VAL_MISSING = "결측"
VAL_NORMAL  = "정상"


# ── 유틸리티 ─────────────────────────────────────────────────────────────────
def load_valid_codes(kan_path: str) -> set:
    """KAN 파일에서 유효한 소분류코드 6자리 집합을 반환합니다."""
    kan = pd.read_excel(kan_path, dtype=str)
    # KAN_CODE는 8자리; 앞 6자리가 소분류 코드. zfill로 선행 0 보장
    codes = kan["KAN_CODE"].dropna().str.strip().str[:6].str.zfill(6)
    return set(codes.unique())


def is_correct(row, col1: str, col2: str, col3: str) -> bool:
    """대/중/소분류 명칭 텍스트 완전 일치 여부."""
    return (
        str(row[ANS_CLS1]).strip() == str(row[col1]).strip()
        and str(row[ANS_CLS2]).strip() == str(row[col2]).strip()
        and str(row[ANS_CLS3]).strip() == str(row[col3]).strip()
    )


def accuracy_on(df: pd.DataFrame, col1: str, col2: str, col3: str,
                mask=None) -> tuple[int, int, float]:
    """mask가 있으면 해당 행만, 없으면 전체에서 정답 수/전체 수/정확도(%) 반환."""
    sub = df if mask is None else df[mask]
    if len(sub) == 0:
        return 0, 0, float("nan")
    correct_mask = sub.apply(lambda r: is_correct(r, col1, col2, col3), axis=1)
    correct = int(correct_mask.sum())
    total   = len(sub)
    return correct, total, round(correct / total * 100, 1)


def code_error_stats(df: pd.DataFrame, code_col: str,
                     valid_codes: set) -> tuple[int, int, float]:
    """소분류코드 컬럼에서 유효하지 않은 코드의 수/전체/비율(%) 반환.

    Excel/pandas가 선행 0을 숫자로 읽어 제거하는 문제를 방지하기 위해
    zfill(6)으로 6자리 문자열로 복원한 뒤 비교합니다.
    """
    codes   = df[code_col].fillna("").astype(str).str.strip().str.zfill(6)
    invalid = ~codes.isin(valid_codes)
    n_err   = int(invalid.sum())
    total   = len(df)
    return n_err, total, round(n_err / total * 100, 1)


# ── 채점 메인 ─────────────────────────────────────────────────────────────────
def score_all(valid_codes: set) -> pd.DataFrame:
    rows = []
    for center, round_, fname in INPUT_FILES:
        fpath = str(_DATA / fname)
        print(f"  로딩: {fname}")
        df = pd.read_excel(fpath, dtype=str)

        missing_mask = df[TEST_CASE].str.strip() == VAL_MISSING
        normal_mask  = df[TEST_CASE].str.strip() == VAL_NORMAL

        for model_name, code_col, cls1, cls2, cls3 in MODELS:
            required = [code_col, cls1, cls2, cls3, ANS_CLS1, ANS_CLS2, ANS_CLS3, TEST_CASE]
            missing_cols = [c for c in required if c not in df.columns]
            if missing_cols:
                print(f"    {model_name}: 컬럼 없음 {missing_cols} → 건너뜀")
                continue

            n_err, n_total, err_pct  = code_error_stats(df, code_col, valid_codes)
            c_all,  t_all,  acc_all  = accuracy_on(df, cls1, cls2, cls3)
            c_mis,  t_mis,  acc_mis  = accuracy_on(df, cls1, cls2, cls3, missing_mask)
            c_nor,  t_nor,  acc_nor  = accuracy_on(df, cls1, cls2, cls3, normal_mask)

            rows.append({
                "센터":          center,
                "회차":          round_,
                "모델":          model_name,
                "전체_행수":     n_total,
                "코드오류_수":   n_err,
                "코드오류_%":    err_pct,
                "전체_정답":     c_all,
                "전체_정확도_%": acc_all,
                "결측_행수":     t_mis,
                "결측_정답":     c_mis,
                "결측_정확도_%": acc_mis,
                "정상_행수":     t_nor,
                "정상_정답":     c_nor,
                "정상_정확도_%": acc_nor,
            })

    return pd.DataFrame(rows)


# ── 요약 집계 ─────────────────────────────────────────────────────────────────
def make_summary(df: pd.DataFrame) -> pd.DataFrame:
    """모델별로 전체 6파일에 걸친 평균 지표를 집계합니다."""
    agg = (
        df.groupby("모델")
        .agg(
            전체_행수합=("전체_행수", "sum"),
            코드오류_수합=("코드오류_수", "sum"),
            전체_정답합=("전체_정답", "sum"),
            결측_행수합=("결측_행수", "sum"),
            결측_정답합=("결측_정답", "sum"),
            정상_행수합=("정상_행수", "sum"),
            정상_정답합=("정상_정답", "sum"),
        )
        .reset_index()
    )
    agg["코드오류_%_평균"]    = (agg["코드오류_수합"]  / agg["전체_행수합"]  * 100).round(1)
    agg["전체_정확도_%_평균"] = (agg["전체_정답합"]   / agg["전체_행수합"]  * 100).round(1)
    agg["결측_정확도_%_평균"] = (agg["결측_정답합"]   / agg["결측_행수합"]  * 100).round(1)
    agg["정상_정확도_%_평균"] = (agg["정상_정답합"]   / agg["정상_행수합"]  * 100).round(1)

    return agg[[
        "모델", "전체_행수합", "코드오류_수합", "코드오류_%_평균",
        "전체_정확도_%_평균", "결측_정확도_%_평균", "정상_정확도_%_평균",
    ]].rename(columns={
        "전체_행수합":       "총 평가행",
        "코드오류_수합":     "코드오류 수",
        "코드오류_%_평균":   "코드오류율(%)",
        "전체_정확도_%_평균": "전체 정확도(%)",
        "결측_정확도_%_평균": "결측 정확도(%)",
        "정상_정확도_%_평균": "정상 정확도(%)",
    })


def make_pivot(df: pd.DataFrame, metric: str) -> pd.DataFrame:
    """센터×회차를 행으로, 모델을 열로 하는 피벗 테이블."""
    pivot = df.pivot_table(
        index=["센터", "회차"], columns="모델", values=metric
    ).reset_index()
    pivot.columns.name = None
    return pivot


# ── Excel 스타일링 ────────────────────────────────────────────────────────────
C = {
    "blue":   "2F5597",
    "orange": "C55A11",
    "green":  "375623",
    "red":    "9C0006",
    "light_green": "E2EFDA",
    "light_yellow": "FFF2CC",
    "light_red":    "FCE4D6",
    "light_blue":   "DEEAF1",
    "light_orange": "FCE4D6",
    "gray":   "F2F2F2",
    "white":  "FFFFFF",
}

def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)

def _font(bold=False, color="000000", size=10) -> Font:
    return Font(bold=bold, color=color, size=size)

def _border() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center")

def _auto_width(ws, min_w=8, max_w=22):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max(
            len(str(col_cells[0].value or "")),
            max((len(str(c.value or "")) for c in col_cells[1:]), default=0),
        )
        ws.column_dimensions[col_letter].width = min(max(max_len * 1.3 + 2, min_w), max_w)

def _style_header_row(ws, fill_color: str):
    for cell in ws[1]:
        cell.fill      = _fill(fill_color)
        cell.font      = _font(bold=True, color="FFFFFF", size=10)
        cell.alignment = _center()
        cell.border    = _border()
    ws.row_dimensions[1].height = 30

def _style_data_rows(ws, pct_cols: set[int]):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font      = _font(size=10)
            cell.alignment = _center()
            cell.border    = _border()

            if cell.column in pct_cols and cell.value is not None:
                try:
                    val = float(cell.value)
                    if val >= 80:
                        cell.fill = _fill(C["light_green"])
                        cell.font = _font(bold=True, color=C["green"])
                    elif val >= 60:
                        cell.fill = _fill(C["light_yellow"])
                    else:
                        cell.fill = _fill(C["light_red"])
                        cell.font = _font(bold=True, color=C["red"])
                except (ValueError, TypeError):
                    pass


def save_excel(detail_df: pd.DataFrame, summary_df: pd.DataFrame) -> None:
    model_order = ["GPT", "Gemini", "Claude", "Clova"]

    # 피벗 테이블 (4종)
    pv_err  = make_pivot(detail_df, "코드오류_%")
    pv_all  = make_pivot(detail_df, "전체_정확도_%")
    pv_mis  = make_pivot(detail_df, "결측_정확도_%")
    pv_nor  = make_pivot(detail_df, "정상_정확도_%")

    # 열 순서 정리 (센터·회차 + 모델 순)
    def reorder(pv):
        fixed = ["센터", "회차"]
        ordered = fixed + [m for m in model_order if m in pv.columns]
        return pv[ordered]

    pv_err = reorder(pv_err)
    pv_all = reorder(pv_all)
    pv_mis = reorder(pv_mis)
    pv_nor = reorder(pv_nor)

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="종합요약",      index=False)
        detail_df .to_excel(writer, sheet_name="상세결과",      index=False)
        pv_all    .to_excel(writer, sheet_name="전체정확도",    index=False)
        pv_mis    .to_excel(writer, sheet_name="결측치카테고리_정확도",    index=False)
        pv_nor    .to_excel(writer, sheet_name="정상카테고리_정확도",    index=False)
        pv_err    .to_excel(writer, sheet_name="코드오류율",    index=False)

    wb = openpyxl.load_workbook(OUTPUT_FILE)

    # ── 종합요약 시트 ────────────────────────────────────────────────────
    ws = wb["종합요약"]
    _style_header_row(ws, C["blue"])
    # 정확도 컬럼 인덱스 찾기
    header = {cell.value: cell.column for cell in ws[1]}
    pct_cols_sum = {
        v for k, v in header.items()
        if k and ("정확도" in str(k) or "코드오류율" in str(k))
    }
    # 코드오류율은 낮을수록 좋으므로 별도 처리
    err_col = header.get("코드오류율(%)")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font      = _font(size=10)
            cell.alignment = _center()
            cell.border    = _border()

            # 모델명 굵게
            if cell.column == 1:
                cell.font = _font(bold=True, size=11)

            # 정확도 열 색상
            if cell.column in pct_cols_sum - ({err_col} if err_col else set()):
                try:
                    val = float(cell.value)
                    if val >= 80:
                        cell.fill = _fill(C["light_green"])
                        cell.font = _font(bold=True, color=C["green"])
                    elif val >= 60:
                        cell.fill = _fill(C["light_yellow"])
                    else:
                        cell.fill = _fill(C["light_red"])
                        cell.font = _font(bold=True, color=C["red"])
                except (ValueError, TypeError):
                    pass

            # 코드오류율: 높을수록 빨간색
            if err_col and cell.column == err_col:
                try:
                    val = float(cell.value)
                    if val == 0:
                        cell.fill = _fill(C["light_green"])
                        cell.font = _font(bold=True, color=C["green"])
                    elif val <= 5:
                        cell.fill = _fill(C["light_yellow"])
                    else:
                        cell.fill = _fill(C["light_red"])
                        cell.font = _font(bold=True, color=C["red"])
                except (ValueError, TypeError):
                    pass

    _auto_width(ws)
    ws.freeze_panes = "B2"

    # ── 상세결과 시트 ───────────────────────────────────────────────────────
    ws2 = wb["상세결과"]
    _style_header_row(ws2, C["blue"])
    header2 = {cell.value: cell.column for cell in ws2[1]}
    pct_cols_det = {
        v for k, v in header2.items()
        if k and "정확도_%" in str(k)
    }
    err_col2 = header2.get("코드오류_%")
    _style_data_rows(ws2, pct_cols_det)
    # 코드오류 열 별도 색상 (낮을수록 좋음)
    if err_col2:
        for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
            cell = row[err_col2 - 1]
            try:
                val = float(cell.value)
                if val == 0:
                    cell.fill = _fill(C["light_green"])
                    cell.font = _font(bold=True, color=C["green"])
                elif val <= 5:
                    cell.fill = _fill(C["light_yellow"])
                else:
                    cell.fill = _fill(C["light_red"])
                    cell.font = _font(bold=True, color=C["red"])
            except (ValueError, TypeError):
                pass
    _auto_width(ws2)
    ws2.freeze_panes = "D2"

    # ── 피벗 시트 공통 스타일 ────────────────────────────────────────────────
    pivot_sheets = {
        "전체정확도":           C["blue"],
        "결측치카테고리_정확도": C["orange"],
        "정상카테고리_정확도":   C["green"],
        "코드오류율":           C["red"],
    }
    err_sheet = "코드오류율"

    for sheet_name, hdr_color in pivot_sheets.items():
        wsp = wb[sheet_name]
        _style_header_row(wsp, hdr_color)
        headerp = {cell.value: cell.column for cell in wsp[1]}
        model_cols = {v for k, v in headerp.items() if k in model_order}

        for row in wsp.iter_rows(min_row=2, max_row=wsp.max_row):
            for cell in row:
                cell.font      = _font(size=10)
                cell.alignment = _center()
                cell.border    = _border()

                if cell.column not in model_cols:
                    cell.font = _font(bold=True, size=10)
                    cell.fill = _fill(C["gray"])
                    continue

                if cell.value is None:
                    continue
                try:
                    val = float(cell.value)
                except (ValueError, TypeError):
                    continue

                if sheet_name == err_sheet:
                    # 오류율: 낮을수록 좋음
                    if val == 0:
                        cell.fill = _fill(C["light_green"])
                        cell.font = _font(bold=True, color=C["green"])
                    elif val <= 5:
                        cell.fill = _fill(C["light_yellow"])
                    else:
                        cell.fill = _fill(C["light_red"])
                        cell.font = _font(bold=True, color=C["red"])
                else:
                    # 정확도: 높을수록 좋음
                    if val >= 80:
                        cell.fill = _fill(C["light_green"])
                        cell.font = _font(bold=True, color=C["green"])
                    elif val >= 60:
                        cell.fill = _fill(C["light_yellow"])
                    else:
                        cell.fill = _fill(C["light_red"])
                        cell.font = _font(bold=True, color=C["red"])

        _auto_width(wsp)
        wsp.freeze_panes = "C2"

    wb.save(OUTPUT_FILE)
    print(f"\n저장 완료 → '{OUTPUT_FILE}'")


# ── 콘솔 출력 ─────────────────────────────────────────────────────────────────
def print_summary(detail_df: pd.DataFrame, summary_df: pd.DataFrame) -> None:
    SEP = "=" * 72

    print(f"\n{SEP}")
    print("  종합 모델 평가 결과 — 센터별 × 회차별 × 모델별")
    print(SEP)

    for center in ["A", "B"]:
        print(f"\n■ {center}센터")
        sub = detail_df[detail_df["센터"] == center]
        pivot = sub.pivot_table(
            index=["회차", "모델"],
            values=["코드오류_%", "전체_정확도_%", "결측_정확도_%", "정상_정확도_%"],
            aggfunc="first"
        )[["코드오류_%", "전체_정확도_%", "결측_정확도_%", "정상_정확도_%"]]
        pivot.columns = ["코드오류율%", "전체정확도%", "결측정확도%", "정상정확도%"]
        print(pivot.to_string(float_format="%.1f"))

    print(f"\n{SEP}")
    print("  모델별 종합 집계 (6파일 전체)")
    print(SEP)
    print(summary_df.to_string(index=False, float_format="%.1f"))
    print()


# ── 실행 진입점 ────────────────────────────────────────────────────────────────
def main():
    import os
    import sys
    import shutil
    sys.stdout.reconfigure(encoding="utf-8")

    if not os.path.exists(KAN_FILE):
        print(f"KAN 코드 파일 없음: '{KAN_FILE}'")
        return

    # 기존 결과 파일 백업
    if os.path.exists(OUTPUT_FILE):
        backup = OUTPUT_FILE.replace(".xlsx", "_backup.xlsx")
        shutil.copy2(OUTPUT_FILE, backup)
        print(f"백업 완료: {backup}\n")

    print("KAN 유효 코드 로딩...")
    valid_codes = load_valid_codes(KAN_FILE)
    print(f"  → 유효 소분류코드 {len(valid_codes):,}개 로딩 완료")

    print("\n채점 시작...")
    detail_df  = score_all(valid_codes)
    summary_df = make_summary(detail_df)

    print_summary(detail_df, summary_df)

    print("엑셀 저장 중...")
    save_excel(detail_df, summary_df)


if __name__ == "__main__":
    main()
